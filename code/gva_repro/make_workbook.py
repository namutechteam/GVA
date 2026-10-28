"""Source-data workbook for every table and figure (Supplementary Data 1)."""
import json

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from facts import compute
from gva.noise import bias_unmitigated

F = compute()
sm_runs = json.load(open("results/small_molecules.json"))["runs"]
cis = F["cis"]
wb = Workbook()
HDR = Font(bold=True)
FILL = PatternFill("solid", fgColor="E7E6E6")
index_rows = []


import numbering


def sheet(name, header, rows, desc, element):
    name, desc, element = numbering.remap_sheet(name), numbering.remap(desc), numbering.remap(element)
    ws = wb.create_sheet(name)
    ws.append(header)
    for c in ws[1]:
        c.font = HDR
        c.fill = FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([float(x) if isinstance(x, (np.floating,)) else
                   int(x) if isinstance(x, (np.integer,)) else x for x in r])
    for j in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(j)].width = max(11, min(34, len(str(header[j - 1])) + 2))
    ws.freeze_panes = "A2"
    index_rows.append((name, element, desc))
    return ws


# ------------------------------------------------------------- Table 1
rows = []
for s in F["small_systems"]:
    key = (s["system"], s["scale"])
    g = F["small_groups"]
    row = [s["system"], s["scale"], s["n_qubits"], s["n_electrons"],
           s["E_FCI"], s["E_corr"] * 1e3, s["c0_sq"],
           (s["E_CISD"] - s["E_FCI"]) * 1e3, (s["E_CCSD"] - s["E_FCI"]) * 1e3,
           (s["E_CCSD(T)"] - s["E_FCI"]) * 1e3,
           s["counts"]["UCCSD"]["n_params"], s["counts"]["GVA"]["n_params"],
           s["counts"]["UCCSD"]["cnot_1L"], s["counts"]["GVA"]["cnot_1L"]]
    for an in ("UCCSD", "GVA"):
        for L in (1, 2):
            e = np.array([r["dE"] for r in g[key + (an, L)]]) * 1e3
            hf = [r["dE"] for r in g[key + (an, L)] if r["seed"] == 0][0] * 1e3
            row += [hf, np.median(e), e.min(), e.max()]
    ad = g[key + ("ADAPT-VQE", None)][0]
    row += [ad["dE"] * 1e3, ad["n_params"], ad["cnot"]]
    rows.append(row)
hdr = ["System", "Bond scale", "Qubits", "Electrons", "E_FCI (Ha)", "E_corr (mHa)",
       "c0^2 (HF weight in FCI)", "CISD error (mHa)", "CCSD error (mHa)",
       "CCSD(T) error (mHa)", "UCCSD params/layer", "GVA params/layer",
       "UCCSD CNOT/layer", "GVA CNOT/layer"]
for an in ("UCCSD", "GVA"):
    for L in (1, 2):
        hdr += [f"{an} {L}L dE HF start (mHa)", f"{an} {L}L dE median 5 starts (mHa)",
                f"{an} {L}L dE min (mHa)", f"{an} {L}L dE max (mHa)"]
hdr += ["ADAPT-VQE dE (mHa)", "ADAPT operators", "ADAPT CNOT"]
sheet("Table1_SmallMolecules", hdr, rows,
      "Classical references, resource counts and VQE errors relative to FCI (STO-3G).",
      "Table 1; Fig. 2a-c")

# ------------------------------------------------------------- all small runs
rows = []
for r in sm_runs:
    rows.append([r["system"], r["scale"], r["ansatz"], r["layers"], r["seed"],
                 r["n_params"], r["cnot"], r["E"], r["dE"] * 1e3, r.get("nit"),
                 r.get("nfev"), r.get("iter_to_chem_acc"), r.get("S2"),
                 r.get("wall_s"), r.get("stop_reason")])
sheet("Runs_SmallMolecules",
      ["System", "Bond scale", "Ansatz", "Layers", "Seed (0 = HF start)", "Parameters",
       "CNOT", "E final (Ha)", "dE vs FCI (mHa)", "L-BFGS iterations", "Function evals",
       "Iteration reaching 1.6 mHa", "<S^2>", "Wall time (s)", "ADAPT stop reason"],
      rows, "Every small-molecule optimization run (no omissions).",
      "Table 1; Fig. 2; Supplementary Fig. 1")

# ------------------------------------------------------------- Fig2d histories
ef = F["sysd"][("BeH2", 2.0)]["E_FCI"]
hist = {}
for an in ("UCCSD", "GVA"):
    for L in (1, 2):
        r = [x for x in F["small_groups"][("BeH2", 2.0, an, L)] if x["seed"] == 0][0]
        hist[f"{an} {L}L"] = (np.array(r["history"]) - ef) * 1e3
n = max(len(v) for v in hist.values())
rows = [[i] + [v[i] if i < len(v) else None for v in hist.values()] for i in range(n)]
sheet("Fig2d_BeH2x2_Convergence", ["Iteration"] + [f"{k} dE (mHa)" for k in hist], rows,
      "Energy error vs L-BFGS iteration, BeH2 bonds x2.0, HF start.", "Fig. 2d")

# ------------------------------------------------------------- Fig3
rows = []
for m in F["meas"]:
    row = [m["system"], m["n_qubits"], m["n_pauli_terms"], m["identity_coeff"],
           m["shots_ungrouped"]]
    for kind in ("QWC", "GC"):
        for meth in ("largest_first", "dsatur", "sorted_insertion"):
            row += [m[kind][meth]["n_groups"], m[kind][meth]["shots"]]
    rows.append(row)
hdr = ["System", "Qubits", "Pauli terms (incl. identity)", "Identity coefficient (Ha)",
       "Shots, no grouping"]
for kind in ("QWC", "GC"):
    for meth in ("largest_first", "dsatur", "sorted_insertion"):
        hdr += [f"{kind} {meth} groups", f"{kind} {meth} shots"]
sheet("Fig3_Measurement", hdr, rows,
      "Measurement groups and optimal-allocation shots for sigma = 1.6 mHa, exact ground state.",
      "Fig. 3; Supplementary Table S6")

# ------------------------------------------------------------- Table2 / cisplatin runs
rows = []
for r in cis["runs"]:
    rows.append([r["ansatz"], r["layers"], r["seed"], r["n_params"], r["cnot"], r["E"],
                 r["dE"] * 1e3, r.get("nit"), r.get("nfev"), r.get("iter_to_chem_acc"),
                 r.get("S2"), r.get("wall_s"), r.get("stop_reason")])
sheet("Runs_Cisplatin",
      ["Ansatz", "Layers", "Seed (0 = HF start)", "Parameters", "CNOT", "E final (Ha)",
       "dE vs CASCI (mHa)", "L-BFGS iterations", "Function evals",
       "Iteration reaching 1.6 mHa", "<S^2>", "Wall time (s)", "ADAPT stop reason"],
      rows, "Every cisplatin CAS(10e,10o) run at the C2v geometry.",
      "Table 2; Fig. 4b,c; Fig. 5; Supplementary Table S3")
ref_rows = [["RHF (reference determinant)", cis["E_ref"], (cis["E_ref"] - cis["E_CASCI"]) * 1e3],
            ["CISD", cis["E_CISD"], (cis["E_CISD"] - cis["E_CASCI"]) * 1e3],
            ["CCSD", cis["E_CCSD"], (cis["E_CCSD"] - cis["E_CASCI"]) * 1e3],
            ["CCSD(T)", cis["E_CCSD(T)"], (cis["E_CCSD(T)"] - cis["E_CASCI"]) * 1e3],
            ["CASCI", cis["E_CASCI"], 0.0]]
sheet("Table2_Cisplatin_References", ["Method (same CAS(10e,10o))", "E (Ha)", "E - E_CASCI (mHa)"],
      ref_rows, "Classical references inside the active space.", "Table 2; Fig. 4b; Fig. 5")

# Fig4 histories
hist = {}
for an in ("GVA", "UCCSD"):
    for L in (1, 2, 3):
        rr = [x for x in cis["runs"] if x["ansatz"] == an and x["layers"] == L and x["seed"] == 0]
        if rr:
            hist[f"{an} {L}L"] = (np.array(rr[0]["history"]) - cis["E_CASCI"]) * 1e3
n = max(len(v) for v in hist.values())
rows = [[i] + [v[i] if i < len(v) else None for v in hist.values()] for i in range(n)]
sheet("Fig4b_Cisplatin_Convergence", ["Iteration"] + [f"{k} dE (mHa)" for k in hist], rows,
      "Energy error vs L-BFGS iteration, HF start.", "Fig. 4b")
ad = [x for x in cis["runs"] if x["ansatz"] == "ADAPT-VQE"]
if ad:
    rows = [[t["n_ops"], t["E"], (t["E"] - cis["E_CASCI"]) * 1e3, t["grad_norm"], t["wall_s"]]
            for t in ad[0]["trace"]]
    sheet("Fig4c_Cisplatin_ADAPT", ["Operators", "E (Ha)", "dE (mHa)", "Pool gradient norm",
                                    "Elapsed (s)"], rows, "ADAPT-VQE trace (GVA pool).", "Fig. 4c")
orb = json.load(open("results/cisplatin_orbitals.json"))["orbitals"]
rows = [[k + 1, o["irrep"], o["occupation"], o["composition"]] for k, o in enumerate(orb)]
sheet("Fig4a_ActiveNOs", ["Active NO", "Irrep (C2v)", "MP2 occupation",
                          "Gross population by element/angular momentum"], rows,
      "Occupation numbers of the ten active MP2 natural orbitals.", "Fig. 4a; Supplementary Table S9")

# ------------------------------------------------------------- Fig6
rows = [[r["n_active_orbitals"], r["n_qubits"], r["UCCSD_params"], r["UCCSD_formula"],
         r["PGonly_params"], r["spinonly_params"], r["GVA_params"], r["UCCSD_cnot_1L"],
         r["GVA_cnot_1L"], r["GVA_params"] / r["UCCSD_params"],
         r["GVA_cnot_1L"] / r["UCCSD_cnot_1L"]] for r in F["scal"]]
sheet("Fig6_Scaling", ["Active orbitals", "Qubits", "UCCSD params", "UCCSD formula",
                       "PG-only params", "Spin-only params", "GVA params", "UCCSD CNOT/layer",
                       "GVA CNOT/layer", "Param ratio GVA/UCCSD", "CNOT ratio GVA/UCCSD"], rows,
      "Resource counts vs active-space size (cisplatin MP2 NOs, C2v).", "Fig. 6")

# ------------------------------------------------------------- Fig7
rows = [[r["system"], r["ansatz"], r["layers"], r["cnot"], r["dE_noiseless_mHa"], r["gap_Ha"],
         r["p_max_unmitigated"], r["p_max_richardson"], r["p_max_rescaling_overhead"]]
        for r in F["noise"]]
sheet("Fig7_NoiseBudget", ["System", "Ansatz", "Layers", "CNOT", "Noiseless dE (mHa)",
                           "c0 - E (Ha)", "p_max unmitigated", "p_max Richardson (1,3,5)",
                           "p_max for 1/F^2 <= 100"], rows,
      "Global depolarizing error budget (closed form; see Methods).", "Fig. 7b")
cm = [m for m in F["meas"] if m["system"].startswith("Cisplatin")][0]
gap = cm["identity_coeff"] - cm["E_exact"]
p = np.logspace(-10, -2, 81)
cols = []
for an in ("GVA", "UCCSD"):
    for L in (1, 2):
        rr = [x for x in cis["runs"] if x["ansatz"] == an and x["layers"] == L and x["seed"] == 0][0]
        cols.append((f"{an} {L}L bias (mHa)", bias_unmitigated(p, rr["cnot"], gap) * 1e3))
rows = [[p[i]] + [c[1][i] for c in cols] for i in range(len(p))]
sheet("Fig7a_Cisplatin_Bias", ["p (two-qubit depolarizing)"] + [c[0] for c in cols], rows,
      "Unmitigated energy bias vs p for cisplatin circuits.", "Fig. 7a")

# ------------------------------------------------------------- Fig8
rows = []
for r in F["pes"]:
    ef = r["E_CASCI"]
    rows.append([r["R_PtCl"], r["point_group"], r["E_ref"], r["E_CISD"], r["E_CCSD"],
                 r["E_CCSD(T)"], ef, r["GVA_1L"]["E"], r["GVA_2L"]["E"], r["UCCSD_1L"]["E"],
                 (r["E_CISD"] - ef) * 1e3, (r["E_CCSD"] - ef) * 1e3, (r["E_CCSD(T)"] - ef) * 1e3,
                 r["GVA_1L"]["dE"] * 1e3, r["GVA_2L"]["dE"] * 1e3, r["UCCSD_1L"]["dE"] * 1e3,
                 r["c0_sq"], r["T1_diag"], r["counts"]["GVA"]["n_params"],
                 r["counts"]["UCCSD"]["n_params"], r["CCSD_converged"]])
sheet("Fig8_PtCl_Scan", ["R(Pt-Cl) (A)", "Point group", "E RHF (Ha)", "E CISD", "E CCSD",
                         "E CCSD(T)", "E CASCI", "E GVA 1L", "E GVA 2L", "E UCCSD 1L",
                         "CISD err (mHa)", "CCSD err (mHa)", "CCSD(T) err (mHa)",
                         "GVA 1L err (mHa)", "GVA 2L err (mHa)", "UCCSD 1L err (mHa)",
                         "c0^2", "T1 diagnostic", "GVA params/layer (Cs)",
                         "UCCSD params/layer", "CCSD converged"], rows,
      "Pt-Cl bond-stretching scan in Cs, CAS(10e,10o).", "Fig. 8; Supplementary Table S7")

# ------------------------------------------------------------- orbital choice
from facts import orbital_sensitivity
rows = []
for r in orbital_sensitivity():
    rows.append([r["choice"], r["selected"], r["E_CASCI"], r["E_ref"], r["E_corr_active_mHa"],
                 r["c0_sq"], r["CISD_err_mHa"], r["CCSD_err_mHa"], r["CCSD(T)_err_mHa"],
                 r["counts"]["UCCSD"]["n_params"], r["counts"]["UCCSD"]["cnot_1L"],
                 r["counts"]["GVA"]["n_params"], r["counts"]["GVA"]["cnot_1L"],
                 r["UCCSD_1L"]["dE_mHa"], r["GVA_1L"]["dE_mHa"], r["GVA_2L"]["dE_mHa"],
                 ", ".join(r["active_irreps"]), " | ".join(r["active_composition"])])
sheet("TableS10_OrbitalChoice", ["Active orbitals", "Selection", "E_CASCI (Ha)", "E_ref (Ha)",
                                 "E_corr active (mHa)", "c0^2", "CISD err (mHa)",
                                 "CCSD err (mHa)", "CCSD(T) err (mHa)", "UCCSD params",
                                 "UCCSD CNOT/layer", "GVA params", "GVA CNOT/layer",
                                 "UCCSD 1L dE (mHa)", "GVA 1L dE (mHa)", "GVA 2L dE (mHa)",
                                 "Active irreps", "Active orbital composition"], rows,
      "Cisplatin CAS(10e,10o): MP2 natural, canonical RHF and CASSCF orbitals.",
      "Main text (orbital choice); Supplementary Table S10")

# ------------------------------------------------------------- final parameters
rows = []
for r in cis["runs"]:
    if r["ansatz"] == "ADAPT-VQE":
        labels = [(1, " + ".join(g)) for g in r["operators"]]
    else:
        pm = cis["parameter_maps_1L"][r["ansatz"]]
        labels = [(L + 1, " + ".join(p["generators"])) for L in range(r["layers"]) for p in pm]
    for k, (th, (layer, gen)) in enumerate(zip(r["theta"], labels)):
        rows.append([r["ansatz"], r["layers"], r["seed"], k, layer, gen, th])
sheet("Theta_Cisplatin", ["Ansatz", "Layers", "Seed", "Parameter index", "Layer",
                          "Generators (tied, applied in this order)", "theta (rad)"], rows,
      "Final variational parameters of every cisplatin run with their generators. "
      "Label 'D 0a,1b->5a,6b' = annihilate spin orbitals 0alpha,1beta; create 5alpha,6beta.",
      "Table 2; Supplementary Table S3; verify_theta.py")
rows = []
sysmap = {(s["system"], s["scale"]): s for s in F["small_systems"]}
for r in sm_runs:
    if r["ansatz"] == "ADAPT-VQE":
        labels = [(1, " + ".join(g)) for g in r["operators"]]
    else:
        pm = sysmap[(r["system"], r["scale"])]["parameter_maps_1L"][r["ansatz"]]
        labels = [(L + 1, " + ".join(p["generators"])) for L in range(r["layers"]) for p in pm]
    for k, (th, (layer, gen)) in enumerate(zip(r["theta"], labels)):
        rows.append([r["system"], r["scale"], r["ansatz"], r["layers"], r["seed"], k, layer,
                     gen, th])
sheet("Theta_SmallMolecules", ["System", "Bond scale", "Ansatz", "Layers", "Seed",
                               "Parameter index", "Layer", "Generators", "theta (rad)"], rows,
      "Final variational parameters of every small-molecule run with their generators.",
      "Table 1; Supplementary Table S4; verify_theta.py")

# ------------------------------------------------------------- geometries
rows = []
for mol, atoms in F["geo"].items():
    for a, c in atoms:
        rows.append([mol, a] + c)
sheet("Geometries", ["Molecule", "Atom", "x (A)", "y (A)", "z (A)"], rows,
      "Cartesian coordinates of the reference geometries.", "Supplementary Table S1")


# ------------------------------------------------------------- periodic systems
import os
from gva.ansatz import uccsd_pool, excitation_label
if os.path.exists("results/periodic.json"):
    per = json.load(open("results/periodic.json"))
    sp_ = json.load(open("results/scaling_periodic.json"))
    PS = {s["system"]: s for s in per["systems"]}
    rows = []
    for s in per["systems"]:
        rows.append([s["system"], s["r"], s["d"], s["a"], s["n_k"], s["n_qubits"], s["n_electrons"],
                     s["E_KRHF_cell"], s["E_KRHF_supercell"], s["E_HF_from_integrals"],
                     s["E_RHF_supercell_gamma"], s["E_FCI"], s["E_FCI_bloch"],
                     s["E_FCI_supercell_gamma"], abs(s["E_FCI_bloch"] - s["E_FCI"]),
                     abs(s["E_FCI_supercell_gamma"] - s["E_FCI"]), s["E_corr"] * 1e3,
                     s["E_corr"] * 1e3 / s["n_k"], s["c0_sq"], s["t1_diag"],
                     (s["E_CISD"] - s["E_FCI"]) * 1e3, (s["E_CCSD"] - s["E_FCI"]) * 1e3,
                     (s["E_CCSD(T)"] - s["E_FCI"]) * 1e3, s["CCSD_converged"],
                     s["homo"], s["lumo"], s["momentum_violation_max"], s["hermiticity_max"],
                     s["time_reversal_max"], 1 - s["FCI_momentum_sector_weights"][s["K_HF"]]])
    sheet("TableS11_PeriodicSystems",
          ["System", "r (A)", "d (A)", "a (A)", "N_k", "Qubits", "Electrons",
           "E_KRHF per cell (Ha)", "E_KRHF x N_k (Ha)", "E_HF from integrals (Ha)",
           "E_RHF explicit supercell (Ha)", "E_FCI real basis (Ha)", "E_FCI Bloch basis (Ha)",
           "E_FCI explicit supercell (Ha)", "|Bloch - real| (Ha)", "|supercell - real| (Ha)",
           "E_corr (mHa)", "E_corr per cell (mHa)", "c0^2", "T1 diagnostic", "CISD err (mHa)",
           "CCSD err (mHa)", "CCSD(T) err (mHa)", "CCSD converged", "HOMO (Ha)", "LUMO (Ha)",
           "max |momentum-violating integral|", "max Hermiticity violation",
           "max time-reversal violation", "FCI weight outside K_HF"], rows,
          "Periodic H2-chain systems: energies, references and integral checks.",
          "Fig. 10; Table 3; Supplementary Table S11")
    rows = []
    for s in per["systems"]:
        for v, c in s["counts"].items():
            rows.append([s["system"], s["n_k"], v, c["n_params"], c.get("n_generators"),
                         c["cnot_1L"], c.get("n_real_orbits"), c.get("n_zero_orbits")])
    sheet("PeriodicCounts_Systems", ["System", "N_k", "Ansatz", "Parameters per layer",
                                     "Generators per layer", "CNOT per layer",
                                     "Orbits with one real parameter", "Orbits removed (t = 0)"],
          rows, "Parameter and CNOT counts per layer for the benchmark chains (enumerated; equal "
          "to the closed forms).", "Table 3; Fig. 10c; Supplementary Table S14")
    rows = []
    for r in per["runs"]:
        rows.append([r["system"], r["n_k"], r["ansatz"], r["layers"], r["seed"], r["E"],
                     r["dE"] * 1e3, r["S2"], r["overlap_sq_FCI"], r["n_params"], r["cnot"],
                     r.get("nit"), r.get("nfev"), r.get("iter_to_chem_acc"), r["wall_s"],
                     r.get("n_ops"), r.get("stop_reason"), r.get("jw_abs_diff"),
                     r.get("max_grad_forbidden"), r.get("message")])
    sheet("TableS14_PeriodicRuns", ["System", "N_k", "Ansatz", "Layers", "Seed", "E (Ha)",
                                    "E - E_FCI (mHa)", "<S^2>", "|<FCI|psi>|^2", "Parameters",
                                    "CNOT", "Iterations", "Evaluations", "Iter. to 1.6 mHa",
                                    "Wall time (s)", "ADAPT operators", "ADAPT stop",
                                    "|E_JW - E| at run time (Ha)",
                                    "max |grad| momentum-violating", "Optimizer message"], rows,
          "Every periodic VQE and ADAPT-VQE run. Seed 0 = Hartree-Fock start; seeds 1, 2 = "
          "N(0, 0.05^2) starts (numpy seed 3000 + seed).",
          "Fig. 10; Supplementary Fig. 2; Supplementary Table S14")
    rows = []
    for s in per["systems"]:
        for an, L in (("UCCSD", 1), ("GVA-k", 1), ("UCCSD", 2), ("GVA-k", 2), ("ADAPT-VQE", None),
                      ("k only", 1), ("k + spin", 1), ("k + TR", 1)):
            rr = [x for x in per["runs"] if x["system"] == s["system"] and x["ansatz"] == an
                  and x["layers"] == L]
            if not rr:
                continue
            e = np.array([x["dE"] for x in rr]) * 1e3
            hf = [x for x in rr if x["seed"] == 0][0]
            rows.append([s["system"], s["r"], s["d"], s["n_k"], an, L, len(rr), hf["dE"] * 1e3,
                         float(e.min()), float(np.median(e)), float(e.max()),
                         max(x["S2"] for x in rr), hf["n_params"], hf["cnot"]])
        rows.append([s["system"], s["r"], s["d"], s["n_k"], "CCSD", None, 1,
                     (s["E_CCSD"] - s["E_FCI"]) * 1e3, None, None, None, None, None, None])
        rows.append([s["system"], s["r"], s["d"], s["n_k"], "CCSD(T)", None, 1,
                     (s["E_CCSD(T)"] - s["E_FCI"]) * 1e3, None, None, None, None, None, None])
    sheet("Fig10_FigS2_Data", ["System", "r (A)", "d (A)", "N_k", "Method", "Layers", "Starts",
                               "Error, HF start (mHa)", "Error min (mHa)", "Error median (mHa)",
                               "Error max (mHa)", "max <S^2> over starts", "Parameters",
                               "CNOT"], rows,
          "Values plotted in Fig. 10 (errors versus N_k, spacing and CNOT count) and "
          "Supplementary Fig. 2 (tying ablation, <S^2>).", "Fig. 10; Supplementary Fig. 2")
    rows = []
    mol_h = {"H2": 8, "H4": 8, "LiH": 4, "BeH2": 8}
    for s in F["small_systems"]:
        if s["scale"] == 1.0:
            c = s["counts"]
            rows.append(["molecule", s["system"], mol_h[s["system"]], c["UCCSD"]["n_params"],
                         c["GVA"]["n_params"], c["GVA"]["n_params"] / c["UCCSD"]["n_params"],
                         c["UCCSD"]["cnot_1L"], c["GVA"]["cnot_1L"],
                         c["GVA"]["cnot_1L"] / c["UCCSD"]["cnot_1L"]])
    for r in F["scal"]:
        rows.append(["cisplatin CAS (C2v)", f"CAS({r['n_active_orbitals']}e,{r['n_active_orbitals']}o)",
                     4, r["UCCSD_params"], r["GVA_params"], r["GVA_params"] / r["UCCSD_params"],
                     r["UCCSD_cnot_1L"], r["GVA_cnot_1L"], r["GVA_cnot_1L"] / r["UCCSD_cnot_1L"]])
    for key, lab in (("one_d", "1D mesh"), ("three_d", "3D mesh"), ("molecular_crystal", "3D mesh, 8+8 bands")):
        for r in sp_[key]:
            rows.append([f"{lab}, {r['o']}+{r['v']} bands", "x".join(map(str, r["mesh"])), r["n_k"],
                         r["UCCSD_params"], r["GVA_k_params"], r["GVA_k_params"] / r["UCCSD_params"],
                         r["UCCSD_cnot_1L"], r["k_cnot_1L"], r["k_cnot_1L"] / r["UCCSD_cnot_1L"]])
    sheet("Fig9_RatioVsGroupOrder", ["Type", "System / mesh", "Group order h", "UCCSD params",
                                     "GVA params", "Param. ratio", "UCCSD CNOT", "GVA CNOT",
                                     "CNOT ratio"], rows,
          "Ratios of GVA to UCCSD counts versus the order of the symmetry group (point group "
          "for molecules, N_k for crystals).", "Fig. 9a,b")
    rows = []
    for key in ("one_d", "three_d", "molecular_crystal"):
        for r in sp_[key]:
            rows.append([key, "x".join(map(str, r["mesh"])), r["n_k"], r["o"], r["v"], r["n_qubits"],
                         r["UCCSD_params"], r["UCCSD_cnot_1L"], r["k_generators"], r["k_only_params"],
                         r["k_spin_params"], r["k_TR_params"], r["GVA_k_params"], r["k_cnot_1L"],
                         r["S_per_spin"], r["D_ss_per_spin"], r["D_ab"], r["n_self"]])
    sheet("TableS13_PeriodicScaling", ["Set", "Mesh", "N_k", "o per cell", "v per cell", "Qubits",
                                       "UCCSD params", "UCCSD CNOT/layer", "Momentum-allowed generators",
                                       "k only params", "k + spin params", "k + TR params",
                                       "GVA-k params", "GVA-k CNOT/layer", "S (per spin)",
                                       "D_ss (per spin)", "D_ab", "n_self"], rows,
          "Closed-form counts for crystals with o occupied and v virtual bands per cell.",
          "Fig. 9c; Supplementary Table S13")
    rows = []
    for r in sp_["enumeration_check"]:
        rows.append(["x".join(map(str, r["mesh"])), r["o"], r["v"], r["pool"],
                     r["UCCSD_params_enum"], r["closed_form"]["UCCSD_params"],
                     r["UCCSD_cnot_enum"], r["closed_form"]["UCCSD_cnot_1L"],
                     r["k only"]["n_params"], r["closed_form"]["k_only_params"],
                     r["k + spin"]["n_params"], r["closed_form"]["k_spin_params"],
                     r["k + TR"]["n_params"], r["closed_form"]["k_TR_params"],
                     r["GVA-k"]["n_params"], r["closed_form"]["GVA_k_params"],
                     r["GVA-k"]["cnot_1L"], r["closed_form"]["k_cnot_1L"],
                     r["GVA-k"]["n_real_orbits"], r["GVA-k"]["n_zero_orbits"], r["agree"]])
    sheet("TableS12_EnumerationCheck", ["Mesh", "o", "v", "Pool size", "UCCSD params (enum)",
                                        "UCCSD params (formula)", "UCCSD CNOT (enum)",
                                        "UCCSD CNOT (formula)", "k only (enum)", "k only (formula)",
                                        "k + spin (enum)", "k + spin (formula)", "k + TR (enum)",
                                        "k + TR (formula)", "GVA-k (enum)", "GVA-k (formula)",
                                        "GVA-k CNOT (enum)", "GVA-k CNOT (formula)",
                                        "Orbits with one real parameter", "Orbits removed",
                                        "All agree"], rows,
          "Closed-form counts against explicit enumeration of the pool and of the tied orbits.",
          "Supplementary Table S12")
    if os.path.exists("results/measurement_periodic.json"):
        rows = []
        for m in json.load(open("results/measurement_periodic.json")):
            rows.append([m["system"], m["basis"], m["n_qubits"], m["n_pauli_terms"], m["one_norm"],
                         m["shots_ungrouped"], m["QWC"]["n_groups"], m["QWC"]["shots"],
                         m["GC"]["n_groups"], m["GC"]["shots"], m["abs_diff_pruned"]])
        sheet("TableS15_PeriodicMeasurement", ["System", "Basis", "Qubits", "Pauli terms",
                                               "1-norm (Ha)", "Shots, no grouping",
                                               "QWC groups", "QWC shots", "GC groups", "GC shots",
                                               "|E_Pauli - E_exact| after 1e-8 pruning (Ha)"],
              rows, "Measurement cost in the Bloch and real supercell bases (sorted insertion, "
              "1.6 mHa target, exact ground state).", "Supplementary Table S15")
    if os.path.exists("results/tying_check_periodic.json"):
        rows = [[t["system"], t["n_amplitudes"], t["max_abs_t"], t["max_dev_time_reversal"],
                 t["rel_dev_time_reversal"], t["max_dev_spin_flip"], t["rel_dev_spin_flip"],
                 t["E_k_only_minus_E_GVA_k_mHa"]]
                for t in json.load(open("results/tying_check_periodic.json"))]
        sheet("Periodic_TyingCheck", ["System", "Amplitudes", "max |t|", "max |t_TR - s t*|",
                                      "relative", "max |t_flip - t|", "relative",
                                      "E(k only) - E(GVA-k), 1L (mHa)"], rows,
              "Untied single-layer optimum: deviation from the time-reversal and spin-flip "
              "relations imposed by GVA-k.", "Supplementary Note 14")
    if os.path.exists("results/verify_periodic.json"):
        v = json.load(open("results/verify_periodic.json"))
        rows = [[x["system"], x["ansatz"], x["layers"], x["seed"], x["n_params"], x["E_stored"],
                 x["E_rebuilt"], x["abs_dE"], x["S2_stored"], x["S2_rebuilt"], x["abs_dS2"],
                 x.get("E_jordan_wigner_exact"), x.get("jw_abs_diff"), x.get("n_pauli_terms_exact"),
                 x["ok"]] for x in v["runs"]]
        sheet("Periodic_VerifyTheta", ["System", "Ansatz", "Layers", "Seed", "Parameters",
                                       "E stored (Ha)", "E rebuilt (Ha)", "|dE| (Ha)", "S2 stored",
                                       "S2 rebuilt", "|dS2|", "E Jordan-Wigner, exact (Ha)",
                                       "|E_JW - E| (Ha)", "Pauli terms (exact)", "OK"], rows,
              "Every periodic final state rebuilt from the stored parameters (verify_periodic.py).",
              "Supplementary Note 12")

    rows = []
    for r in per["runs"]:
        ef = PS[r["system"]]["E_FCI"]
        if r["ansatz"] == "ADAPT-VQE":
            for t in r.get("trace", []):
                rows.append([r["system"], r["ansatz"], None, r["seed"], t["n_ops"], t["E"],
                             (t["E"] - ef) * 1e3, t["grad_norm"]])
        else:
            for k, e in enumerate(r.get("history", [])):
                rows.append([r["system"], r["ansatz"], r["layers"], r["seed"], k, e,
                             (e - ef) * 1e3, None])
    sheet("Periodic_EnergyHistory", ["System", "Ansatz", "Layers", "Seed",
                                     "Iteration (ADAPT: operators)", "E (Ha)", "E - E_FCI (mHa)",
                                     "ADAPT pool-gradient norm"], rows,
          "Energy after every optimizer iteration of every periodic run (ADAPT-VQE: after "
          "each added orbit).", "Supplementary Table S14")
    rows = []
    pools = {}
    for r in per["runs"]:
        s = PS[r["system"]]
        if r["ansatz"] == "UCCSD":
            key = r["system"]
            if key not in pools:
                pools[key] = [excitation_label(e) for e in uccsd_pool(s["n_electrons"] // 2, s["norb"])]
            labels = [(L + 1, g, "real") for L in range(r["layers"]) for g in pools[key]]
            comps = [(L, g, m, "x") for (L, g, m) in labels]
        else:
            pl = r.get("parameter_labels") or s["parameter_maps_1L"][r["ansatz"]]
            if r["ansatz"] != "ADAPT-VQE" and r["layers"] == 2 and not r.get("parameter_labels"):
                pl = [dict(o, layer=L + 1) for L in range(2) for o in s["parameter_maps_1L"][r["ansatz"]]]
            comps = []
            for o in pl:
                mem = " + ".join(f"{'' if sg > 0 else '-'}{'conj ' if c else ''}{lab}"
                                 for lab, sg, c in o["members"])
                if o["mode"] == "complex":
                    comps += [(o["layer"], mem, o["mode"], "Re t"), (o["layer"], mem, o["mode"], "Im t")]
                elif o["mode"] == "real":
                    comps.append((o["layer"], mem, o["mode"], "Re t"))
                else:
                    comps.append((o["layer"], mem, o["mode"], "Im t"))
        assert len(comps) == len(r["theta"]), (r["system"], r["ansatz"], len(comps), len(r["theta"]))
        for k, (th, (layer, gen, mode, comp)) in enumerate(zip(r["theta"], comps)):
            rows.append([r["system"], r["ansatz"], r["layers"], r["seed"], k, layer, comp, mode,
                         gen, th])
    sheet("Theta_Periodic", ["System", "Ansatz", "Layers", "Seed", "Parameter index", "Layer",
                             "Component", "Orbit mode", "Excitations (sign, conj = amplitude t*)",
                             "Value"], rows,
          "Final parameters of every periodic run. For Bloch-basis ansaetze each orbit carries "
          "t = Re t + i Im t; a member applies sign*t or sign*conj(t). UCCSD (real basis): one "
          "real amplitude per excitation.", "Supplementary Table S14; verify_periodic.py")

# ------------------------------------------------------------- live formulas
ws = wb.create_sheet("ParameterFormulas")
ws.append(["n_o (occupied)", "n_v (virtual)", "N_UCCSD", "N_spin-tied", "CNOT single (p<q)",
           "Note"])
for c in ws[1]:
    c.font = HDR
    c.fill = FILL
for i, (no, nv) in enumerate([(1, 1), (2, 2), (2, 4), (3, 4), (5, 5), (10, 10)], start=2):
    ws.append([no, nv, f"=2*A{i}*B{i}+2*COMBIN(A{i},2)*COMBIN(B{i},2)+(A{i}*B{i})^2",
               f"=A{i}*B{i}+COMBIN(A{i},2)*COMBIN(B{i},2)+A{i}*B{i}*(A{i}*B{i}+1)/2",
               None, "edit n_o, n_v to recompute"])
ws.append([])
ws.append(["CNOT, single excitation with spin orbitals p<q: 4(q-p); double with sorted p<q<r<s: 16(q-p+s-r+1)"])
for j in range(1, 7):
    ws.column_dimensions[get_column_letter(j)].width = 22
index_rows.append(("ParameterFormulas", "Methods eq. (2)-(5)", "Live Excel formulas for the closed-form counts."))

# ------------------------------------------------------------- index
ws = wb["Sheet"]
ws.title = "Index"
ws.append(["Sheet", "Manuscript element", "Content"])
for c in ws[1]:
    c.font = HDR
    c.fill = FILL
for r in index_rows:
    ws.append(list(r))
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 90
wb.save("GVA_Source_Data.xlsx")
print("saved", len(index_rows), "sheets")
