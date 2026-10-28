"""Minimal python-docx helpers for a journal-style manuscript.

Inline markup in strings:  _{sub}  ^{sup}  *italic*  **bold**
"""
import datetime
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm, RGBColor

TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|_\{.+?\}|\^\{.+?\})")
FONT = "Times New Roman"


def _set_font(run, size=None, bold=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


MINUS = re.compile(r"(?<![\w.)\]])-(?=\d)")


def add_runs(p, text, size=None, bold=False, italic=False, sub=False, sup=False):
    """Recursive inline markup: **bold**, *italic*, _{sub}, ^{sup} (one level of
    braces). Hyphen-minus before a digit is typeset as a true minus sign."""
    text = MINUS.sub("−", text)
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            add_runs(p, part[2:-2], size, True, italic, sub, sup)
            continue
        if part.startswith("*") and part.endswith("*") and len(part) > 2 and not part.startswith("**"):
            add_runs(p, part[1:-1], size, bold, True, sub, sup)
            continue
        if part.startswith("_{") and part.endswith("}"):
            add_runs(p, part[2:-1], size, bold, italic, True, False)
            continue
        if part.startswith("^{") and part.endswith("}"):
            add_runs(p, part[2:-1], size, bold, italic, False, True)
            continue
        r = p.add_run(part)
        _set_font(r, size, bold, italic)
        if sub:
            r.font.subscript = True
        if sup:
            r.font.superscript = True
    return p


class Manuscript:
    def __init__(self, title_meta, author="", line_numbers=True, spacing=1.5,
                 size=11.5):
        self.doc = Document()
        self.size = size
        self.spacing = spacing
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.line_spacing = spacing
        pf.space_after = Pt(4)
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(25)
        sec.top_margin = sec.bottom_margin = Mm(22)
        if line_numbers:
            ln = OxmlElement("w:lnNumType")
            ln.set(qn("w:countBy"), "1")
            ln.set(qn("w:restart"), "continuous")
            ln.set(qn("w:distance"), "283")
            sec._sectPr.append(ln)
        self._page_numbers(sec)
        cp = self.doc.core_properties
        cp.title = title_meta
        cp.author = author
        cp.last_modified_by = author
        cp.comments = ""
        cp.keywords = ""
        now = datetime.datetime.utcnow()
        cp.created = now
        cp.modified = now
        cp.revision = 1
        self.remap = lambda t: t          # text hook (figure renumbering)

    def _page_numbers(self, sec):
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = txt
            r._r.append(el)
        _set_font(r, 10)

    # ---------------------------------------------------------------- text
    def title(self, text, size=16):
        p = self.doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(10)
        add_runs(p, self.remap(text), size=size, bold=True)
        return p

    def para(self, text, size=None, align=None, italic=False, bold=False,
             space_after=4, indent_first=False, keep_next=False):
        p = self.doc.add_paragraph()
        add_runs(p, self.remap(text), size=size or self.size, italic=italic, bold=bold)
        p.paragraph_format.space_after = Pt(space_after)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "justify":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if keep_next:
            p.paragraph_format.keep_with_next = True
        return p

    def heading(self, text, level=1):
        sizes = {1: 13, 2: 11.5, 3: 11.5}
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        add_runs(p, self.remap(text), size=sizes[level], bold=True, italic=(level == 3))
        return p

    def equation(self, text, number=None):
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = t.rows[0].cells
        c0.width, c1.width = Mm(140), Mm(20)
        p = c0.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, text, size=self.size)
        if number is not None:
            q = c1.paragraphs[0]
            q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_runs(q, f"({number})", size=self.size)
        return t

    def figure(self, path, caption_title, caption_body, width_mm=165):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=Mm(width_mm))
        p.paragraph_format.keep_with_next = True
        c = self.doc.add_paragraph()
        c.paragraph_format.line_spacing = 1.15
        add_runs(c, self.remap(caption_title) + " ", size=self.size - 1.5, bold=True)
        add_runs(c, self.remap(caption_body), size=self.size - 1.5)
        self.legends = getattr(self, "legends", [])
        self.legends.append((self.remap(caption_title), self.remap(caption_body)))
        c.paragraph_format.space_after = Pt(10)

    def table(self, title, header, rows, footnote=None, col_widths=None,
              size=None):
        size = size or self.size - 2.5
        cap = self.doc.add_paragraph()
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.line_spacing = 1.15
        add_runs(cap, self.remap(title), size=self.size - 1.5)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, h in enumerate(header):
            cell = t.rows[0].cells[j]
            pp = cell.paragraphs[0]
            pp.paragraph_format.line_spacing = 1.0
            add_runs(pp, self.remap(h), size=size, bold=True)
            self._shade(cell, "E7E6E6")
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                pp = t.rows[i + 1].cells[j].paragraphs[0]
                pp.paragraph_format.line_spacing = 1.0
                pp.paragraph_format.space_after = Pt(0)
                add_runs(pp, self.remap(str(v)), size=size)
                if j > 0:
                    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if col_widths:
            t.autofit = False
            tblPr = t._tbl.tblPr
            layout = OxmlElement("w:tblLayout")
            layout.set(qn("w:type"), "fixed")
            tblPr.append(layout)
            grid = t._tbl.tblGrid
            for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
                gc.set(qn("w:w"), str(int(col_widths[j] / 25.4 * 1440)))
            for j, w in enumerate(col_widths):
                for r in t.rows:
                    r.cells[j].width = Mm(w)
        if footnote:
            f = self.doc.add_paragraph()
            f.paragraph_format.line_spacing = 1.15
            add_runs(f, self.remap(footnote), size=self.size - 2.5)
            f.paragraph_format.space_after = Pt(10)
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    @staticmethod
    def _shade(cell, color):
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tcPr.append(shd)

    def references(self, refs):
        for k, r in enumerate(refs, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Mm(7)
            p.paragraph_format.first_line_indent = Mm(-7)
            add_runs(p, f"{k}.\t" + r, size=self.size - 1.5)

    def page_break(self):
        self.doc.add_page_break()

    def save(self, path):
        self.doc.save(path)
